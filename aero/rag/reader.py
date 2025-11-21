"""
Document reader for Aero Agent RAG system.

Handles reading and parsing various document formats:
- Text files (.txt, .md, .rst, .csv, .json, .yaml)
- PDF files (.pdf) - with OCR fallback for scanned PDFs
- Word documents (.docx)
- HTML files (.html, .htm)
- Images (.jpg, .png, .tiff, .bmp) - via OCR
"""

import logging
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Optional
from dataclasses import dataclass

logger = logging.getLogger(__name__)


@dataclass
class ParsedDocument:
    """Represents a parsed document."""

    content: str
    metadata: dict
    source: str
    format: str
    pages: Optional[int] = None
    ocr_used: bool = False


class BaseReader(ABC):
    """Base class for document readers."""

    @property
    @abstractmethod
    def supported_formats(self) -> list[str]:
        """List of supported file extensions."""
        raise NotImplementedError

    @abstractmethod
    def read(self, path: Path) -> ParsedDocument:
        """Read and parse a document."""
        raise NotImplementedError

    def can_read(self, path: Path) -> bool:
        """Check if this reader can handle the file."""
        return path.suffix.lower() in self.supported_formats


class TextReader(BaseReader):
    """Reader for plain text files."""

    @property
    def supported_formats(self) -> list[str]:
        return [".txt", ".md", ".rst", ".csv", ".json", ".yaml", ".yml"]

    def read(self, path: Path) -> ParsedDocument:
        try:
            with open(path, "r", encoding="utf-8") as f:
                content = f.read()
        except UnicodeDecodeError:
            with open(path, "r", encoding="latin-1") as f:
                content = f.read()

        return ParsedDocument(
            content=content,
            metadata={"filename": path.name, "size": path.stat().st_size},
            source=str(path),
            format=path.suffix.lower(),
        )


class PDFReader(BaseReader):
    """Reader for PDF files with OCR fallback."""

    def __init__(self):
        self._backend: Optional[str] = None
        try:
            import pypdf
            self._backend = "pypdf"
        except ImportError:
            try:
                import PyPDF2
                self._backend = "pypdf2"
            except ImportError:
                try:
                    import pdfplumber
                    self._backend = "pdfplumber"
                except ImportError:
                    logger.warning("No PDF backend available")

    @property
    def supported_formats(self) -> list[str]:
        return [".pdf"]

    def read(self, path: Path) -> ParsedDocument:
        if self._backend == "pypdf":
            return self._read_pypdf(path)
        elif self._backend == "pypdf2":
            return self._read_pypdf2(path)
        elif self._backend == "pdfplumber":
            return self._read_pdfplumber(path)
        else:
            # Try OCR as last resort
            return self._read_with_ocr(path)

    def _read_pypdf(self, path: Path) -> ParsedDocument:
        import pypdf
        text_parts = []
        with open(path, "rb") as f:
            reader = pypdf.PdfReader(f)
            page_count = len(reader.pages)
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)

        content = "\n\n".join(text_parts)
        # If no text, might be scanned - try OCR
        if not content.strip():
            return self._read_with_ocr(path)

        return ParsedDocument(
            content=content,
            metadata={"filename": path.name, "backend": "pypdf"},
            source=str(path),
            format=".pdf",
            pages=page_count,
        )

    def _read_pypdf2(self, path: Path) -> ParsedDocument:
        import PyPDF2
        text_parts = []
        with open(path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            page_count = len(reader.pages)
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)

        content = "\n\n".join(text_parts)
        if not content.strip():
            return self._read_with_ocr(path)

        return ParsedDocument(
            content=content,
            metadata={"filename": path.name, "backend": "pypdf2"},
            source=str(path),
            format=".pdf",
            pages=page_count,
        )

    def _read_pdfplumber(self, path: Path) -> ParsedDocument:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(path) as pdf:
            page_count = len(pdf.pages)
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)

        content = "\n\n".join(text_parts)
        if not content.strip():
            return self._read_with_ocr(path)

        return ParsedDocument(
            content=content,
            metadata={"filename": path.name, "backend": "pdfplumber"},
            source=str(path),
            format=".pdf",
            pages=page_count,
        )

    def _read_with_ocr(self, path: Path) -> ParsedDocument:
        try:
            from aero.ocr import get_ocr_text
            content = get_ocr_text(str(path))
            return ParsedDocument(
                content=content,
                metadata={"filename": path.name, "backend": "ocr"},
                source=str(path),
                format=".pdf",
                ocr_used=True,
            )
        except Exception as e:
            logger.warning(f"OCR failed for PDF: {e}")
            return ParsedDocument(
                content=f"[Could not extract text from PDF: {path.name}]",
                metadata={"filename": path.name, "error": str(e)},
                source=str(path),
                format=".pdf",
            )


class DocxReader(BaseReader):
    """Reader for Word documents."""

    def __init__(self):
        self._available = False
        try:
            import docx
            self._available = True
        except ImportError:
            logger.warning("python-docx not installed")

    @property
    def supported_formats(self) -> list[str]:
        return [".docx"]

    def read(self, path: Path) -> ParsedDocument:
        if not self._available:
            return ParsedDocument(
                content=f"[Cannot read .docx - python-docx not installed]",
                metadata={"filename": path.name, "error": "no_backend"},
                source=str(path),
                format=".docx",
            )

        import docx
        doc = docx.Document(path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        content = "\n\n".join(paragraphs)

        return ParsedDocument(
            content=content,
            metadata={"filename": path.name, "paragraphs": len(paragraphs)},
            source=str(path),
            format=".docx",
        )


class HTMLReader(BaseReader):
    """Reader for HTML files."""

    @property
    def supported_formats(self) -> list[str]:
        return [".html", ".htm"]

    def read(self, path: Path) -> ParsedDocument:
        with open(path, "r", encoding="utf-8") as f:
            html = f.read()

        # Try BeautifulSoup first
        try:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(html, "lxml")
            # Remove script and style elements
            for element in soup(["script", "style", "nav", "footer", "header"]):
                element.decompose()
            content = soup.get_text(separator="\n", strip=True)
            title = soup.title.string if soup.title else None
        except ImportError:
            # Fallback to regex-based extraction
            import re
            content = re.sub(r"<script[^>]*>.*?</script>", "", html, flags=re.DOTALL | re.IGNORECASE)
            content = re.sub(r"<style[^>]*>.*?</style>", "", content, flags=re.DOTALL | re.IGNORECASE)
            content = re.sub(r"<[^>]+>", " ", content)
            content = re.sub(r"\s+", " ", content).strip()
            title = None

        return ParsedDocument(
            content=content,
            metadata={"filename": path.name, "title": title},
            source=str(path),
            format=path.suffix.lower(),
        )


class ImageReader(BaseReader):
    """Reader for images via OCR."""

    @property
    def supported_formats(self) -> list[str]:
        return [".jpg", ".jpeg", ".png", ".tiff", ".tif", ".bmp", ".gif"]

    def read(self, path: Path) -> ParsedDocument:
        try:
            from aero.ocr import get_ocr_text
            content = get_ocr_text(str(path))
            return ParsedDocument(
                content=content,
                metadata={"filename": path.name, "backend": "ocr"},
                source=str(path),
                format=path.suffix.lower(),
                ocr_used=True,
            )
        except Exception as e:
            logger.warning(f"OCR failed for image: {e}")
            return ParsedDocument(
                content=f"[Could not extract text from image: {path.name}]",
                metadata={"filename": path.name, "error": str(e)},
                source=str(path),
                format=path.suffix.lower(),
            )


class DocumentReader:
    """
    Unified document reader that handles multiple formats.

    Automatically selects appropriate reader based on file extension.
    Supports: .txt, .md, .pdf, .docx, .html, .jpg, .png, and more.

    Example:
        reader = DocumentReader()
        doc = reader.read("document.pdf")
        print(doc.content)
    """

    def __init__(self, enable_ocr: bool = True):
        """Initialize with all available readers."""
        self._readers: list[BaseReader] = [
            TextReader(),
            PDFReader(),
            DocxReader(),
            HTMLReader(),
        ]
        if enable_ocr:
            self._readers.append(ImageReader())

    def read(self, path: str | Path) -> ParsedDocument:
        """Read a document from file."""
        path = Path(path)

        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

        for reader in self._readers:
            if reader.can_read(path):
                logger.debug(f"Reading {path} with {reader.__class__.__name__}")
                try:
                    return reader.read(path)
                except Exception as e:
                    logger.error(f"Reader {reader.__class__.__name__} failed: {e}")
                    continue

        raise ValueError(f"Unsupported file format: {path.suffix}")

    def read_text(self, text: str, source: str = "inline") -> ParsedDocument:
        """Create a ParsedDocument from raw text."""
        return ParsedDocument(
            content=text,
            metadata={"source_type": "inline"},
            source=source,
            format="text",
        )

    def supported_formats(self) -> list[str]:
        """Get list of all supported file formats."""
        formats = []
        for reader in self._readers:
            formats.extend(reader.supported_formats)
        return formats

    def add_reader(self, reader: BaseReader) -> None:
        """Add a custom reader."""
        self._readers.insert(0, reader)  # Prioritize custom readers
        logger.info(f"Added reader for formats: {reader.supported_formats}")


def read_document(path: str | Path) -> ParsedDocument:
    """Convenience function to read a document."""
    return DocumentReader().read(path)
